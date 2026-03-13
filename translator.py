import io
import os
from concurrent.futures import ThreadPoolExecutor
from typing import Callable, Optional

import anthropic
import docx as python_docx
import requests
from dotenv import load_dotenv

from credits_logger import CreditsLogger

load_dotenv()

LOG_FILE = os.environ.get("LOG_FILE", "/app/logs/translator.log")
os.makedirs(os.path.dirname(LOG_FILE), exist_ok=True)

log = CreditsLogger(LOG_FILE, max_lines=10000, name="translator", console=False)
log.info("translator module loaded — log file: %s", LOG_FILE)

TRANSLATOR_ENDPOINT = os.environ["AZURE_TRANSLATOR_ENDPOINT"].rstrip("/")
TRANSLATOR_KEY = os.environ["AZURE_TRANSLATOR_KEY"]
ANTHROPIC_API_KEY = os.environ["ANTHROPIC_API_KEY"]
ANTHROPIC_BASE_URL = os.environ["ANTHROPIC_BASE_URL"]

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
TRANSLATE_URL = f"{TRANSLATOR_ENDPOINT}/translator/document:translate"
API_VERSION = "2024-05-01"


def translate_docx(
    file_bytes: bytes,
    filename: str,
    source_lang: Optional[str] = None,
    on_status: Optional[Callable[[str], None]] = None,
    on_azure_done: Optional[Callable[[bytes], None]] = None,
    on_claude_done: Optional[Callable[[str], None]] = None,
) -> tuple[bytes, str]:
    """
    Returns (docx_bytes, txt_content):
      - docx_bytes: Azure-translated DOCX with formatting preserved
      - txt_content: Claude plain-text translation of the source document
    """
    def status(msg: str) -> None:
        log.info(msg)
        if on_status:
            on_status(msg)

    def _run_azure() -> bytes:
        params = {"api-version": API_VERSION, "targetLanguage": "en"}
        if source_lang:
            params["sourceLanguage"] = source_lang
        response = requests.post(
            TRANSLATE_URL,
            params=params,
            headers={"Ocp-Apim-Subscription-Key": TRANSLATOR_KEY},
            files={"document": (filename, file_bytes, DOCX_MIME)},
            timeout=300,
        )
        if not response.ok:
            log.error("Azure translation error %s: %s", response.status_code, response.text)
            raise RuntimeError(f"Azure translation failed ({response.status_code}): {response.text}")
        log.info("Azure translation complete. Received %d bytes.", len(response.content))
        if on_azure_done:
            on_azure_done(response.content)
        return response.content

    def _run_claude() -> str:
        source_text = _extract_text(file_bytes)
        log.info("Extracted %d characters of source text.", len(source_text))
        result = _claude_translate(source_text, source_lang)
        log.info("Claude translation complete (%d characters).", len(result))
        if on_claude_done:
            on_claude_done(result)
        return result

    status("Sending to Azure + generating Claude translation...")
    pool = ThreadPoolExecutor(max_workers=2)
    azure_future = pool.submit(_run_azure)
    claude_future = pool.submit(_run_claude)
    pool.shutdown(wait=False)

    translated_bytes = azure_future.result()  # raises if Azure failed
    status("Azure done. Waiting for Claude translation...")

    try:
        txt_content = claude_future.result(timeout=120)
    except Exception as exc:
        log.warning("Claude translation failed — .txt will be empty: %s", exc)
        txt_content = ""
        if on_claude_done:
            on_claude_done(txt_content)

    status("Done.")
    return translated_bytes, txt_content


def _extract_text(docx_bytes: bytes) -> str:
    doc = python_docx.Document(io.BytesIO(docx_bytes))
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    for table in doc.tables:
        for row in table.rows:
            for tc in row._tr.findall(f"{W}tc"):
                for p_elem in tc.findall(f"{W}p"):
                    text = "".join(el.text or "" for el in p_elem.iter(f"{W}t")).strip()
                    if text:
                        parts.append(text)
    return "\n\n".join(parts)


def _claude_translate(text: str, source_lang: Optional[str]) -> str:
    lang_hint = {"nb": "Norwegian Bokmål", "nn": "Norwegian Nynorsk"}.get(source_lang or "", "Norwegian")
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY, base_url=ANTHROPIC_BASE_URL)
    message = client.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=8096,
        messages=[{
            "role": "user",
            "content": (
                f"Translate the following {lang_hint} document to English. "
                "Return only the translated text. "
                "Preserve paragraph breaks exactly as in the original.\n\n"
                f"{text}"
            ),
        }],
    )
    return message.content[0].text
